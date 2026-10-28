"""Парсер каталога товаров с выгрузкой в Excel.

Демонстрация на books.toscrape.com — учебном сайте, созданном специально для парсинга.
Обходит все страницы каталога, собирает название, цену, рейтинг, наличие и ссылку,
сохраняет в оформленный Excel-файл с автофильтром.

Запуск:  python parser.py --pages 5 --out books.xlsx
"""
import argparse
import time
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE_URL = "https://books.toscrape.com/catalogue/page-{}.html"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
HEADERS = {"User-Agent": "Mozilla/5.0 (portfolio parser demo)"}


def fetch(session, url, retries=3):
    for attempt in range(retries):
        try:
            r = session.get(url, headers=HEADERS, timeout=15)
            if r.status_code == 404:
                return None
            r.raise_for_status()
            r.encoding = "utf-8"
            return r.text
        except requests.RequestException:
            if attempt == retries - 1:
                raise
            time.sleep(2 ** attempt)


def parse_page(html, page_url):
    soup = BeautifulSoup(html, "html.parser")
    items = []
    for card in soup.select("article.product_pod"):
        link = card.h3.a
        price = card.select_one(".price_color").text
        rating_cls = card.select_one(".star-rating")["class"][1]
        items.append({
            "title": link["title"],
            "price": float(price.replace("£", "").replace("Â", "")),
            "rating": RATINGS.get(rating_cls, 0),
            "in_stock": "In stock" in card.select_one(".availability").text,
            "url": urljoin(page_url, link["href"]),
        })
    return items


def save_excel(items, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    headers = ["Название", "Цена, £", "Рейтинг", "В наличии", "Ссылка"]
    ws.append(headers)
    for it in items:
        ws.append([it["title"], it["price"], it["rating"], "Да" if it["in_stock"] else "Нет", it["url"]])

    head_fill = PatternFill("solid", fgColor="2F5597")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
    for col, width in zip("ABCDE", (60, 10, 10, 12, 70)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description="Парсер каталога в Excel")
    ap.add_argument("--pages", type=int, default=50, help="сколько страниц обойти (макс. 50)")
    ap.add_argument("--out", default="books.xlsx", help="имя выходного файла")
    ap.add_argument("--delay", type=float, default=0.5, help="пауза между запросами, сек")
    args = ap.parse_args()

    session = requests.Session()
    items = []
    for page in range(1, args.pages + 1):
        url = BASE_URL.format(page)
        html = fetch(session, url)
        if html is None:
            break
        page_items = parse_page(html, url)
        items.extend(page_items)
        print(f"Страница {page}: {len(page_items)} товаров")
        time.sleep(args.delay)

    save_excel(items, args.out)
    print(f"Готово: {len(items)} товаров сохранено в {args.out}")


if __name__ == "__main__":
    main()
