"""Excel-отчёт: «Сводка» с графиком, «Товары», «Изменения цен»."""
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill

ACCENT = "2F5597"
GREEN = PatternFill("solid", fgColor="D9F2E3")
RED = PatternFill("solid", fgColor="FBE0E0")


def _header(ws, row):
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths):
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w


def summary_sheet(wb, items, changes, demo):
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = "Мониторинг цен — отчёт"
    ws["A1"].font = Font(bold=True, size=16, color=ACCENT)
    ws["A2"] = f"Сформирован {datetime.now():%d.%m.%Y %H:%M}" + ("  ·  ДЕМО: изменения цен смоделированы" if demo else "")
    ws["A2"].font = Font(italic=True, color="808080")

    in_stock = sum(it["in_stock"] for it in items)
    cheaper = sum(c["diff"] < 0 for c in changes)
    kpis = [
        ("Товаров в каталоге", len(items)),
        ("Категорий", len({it["category"] for it in items})),
        ("Средняя цена, £", round(sum(it["price"] for it in items) / max(len(items), 1), 2)),
        ("В наличии", f"{in_stock / max(len(items), 1):.0%}"),
        ("Изменений цен", len(changes)),
        ("Подешевело / подорожало", f"{cheaper} / {len(changes) - cheaper}"),
    ]
    for i, (k, v) in enumerate(kpis, start=4):
        ws.cell(i, 1, k).font = Font(color="606060")
        ws.cell(i, 2, v).font = Font(bold=True, size=12)

    by_cat = defaultdict(list)
    for it in items:
        by_cat[it["category"]].append(it["price"])
    start = 4 + len(kpis) + 2
    ws.cell(start, 1, "Категория")
    for col, title in enumerate(["Товаров", "Средняя, £", "Мин., £", "Макс., £"], start=2):
        ws.cell(start, col, title)
    _header(ws, start)
    cats = sorted(by_cat.items(), key=lambda kv: -len(kv[1]))
    for r, (cat, prices) in enumerate(cats, start=start + 1):
        ws.cell(r, 1, cat)
        ws.cell(r, 2, len(prices))
        ws.cell(r, 3, round(sum(prices) / len(prices), 2))
        ws.cell(r, 4, min(prices))
        ws.cell(r, 5, max(prices))
    _widths(ws, (28, 14, 13, 11, 11))

    top = min(len(cats), 12)
    if top:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Средняя цена в крупнейших категориях, £"
        chart.style = 10
        chart.legend = None
        chart.add_data(Reference(ws, min_col=3, min_row=start, max_row=start + top), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=start + 1, max_row=start + top))
        chart.height, chart.width = 9, 16
        ws.add_chart(chart, "G4")


def items_sheet(wb, items):
    ws = wb.create_sheet("Товары")
    ws.append(["Название", "Категория", "Цена, £", "Рейтинг", "В наличии", "Ссылка"])
    _header(ws, 1)
    for it in sorted(items, key=lambda x: (x["category"], x["title"])):
        ws.append([it["title"], it["category"], it["price"], "★" * it["rating"] + "☆" * (5 - it["rating"]),
                   "Да" if it["in_stock"] else "Нет", it["url"]])
    n = ws.max_row
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "0.00"
    ws.conditional_formatting.add(
        f"C2:C{n}", ColorScaleRule(start_type="min", start_color="D9F2E3", mid_type="percentile",
                                   mid_value=50, mid_color="FFFFFF", end_type="max", end_color="FBE0E0"))
    _widths(ws, (55, 20, 10, 10, 11, 60))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def changes_sheet(wb, changes):
    ws = wb.create_sheet("Изменения цен")
    ws.append(["Название", "Категория", "Было, £", "Стало, £", "Разница, £", "Разница, %", "Ссылка"])
    _header(ws, 1)
    if not changes:
        ws.append(["Изменений с прошлого запуска нет"])
    for c in changes:
        ws.append([c["title"], c["category"], c["old_price"], c["price"], c["diff"],
                   c["diff"] / c["old_price"], c["url"]])
        fill = GREEN if c["diff"] < 0 else RED
        for cell in ws[ws.max_row][:6]:
            cell.fill = fill
        ws.cell(ws.max_row, 6).number_format = "+0.0%;-0.0%"
        ws.cell(ws.max_row, 5).number_format = "+0.00;-0.00"
    _widths(ws, (55, 20, 10, 10, 12, 12, 60))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build(items, changes, path, demo=False):
    wb = Workbook()
    summary_sheet(wb, items, changes, demo)
    items_sheet(wb, items)
    changes_sheet(wb, changes)
    wb.save(path)
