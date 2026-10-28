"""Панель администратора: /admin. Доступна только ADMIN_ID из .env."""
from collections import Counter
from datetime import date, timedelta
from io import BytesIO

from aiogram import F, Router
from aiogram.filters import Command
from aiogram.types import BufferedInputFile, CallbackQuery, Message
from aiogram.utils.keyboard import InlineKeyboardBuilder
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import db
from config import ADMIN_ID, MASTERS, SERVICES
from keyboards import fmt_day, plural

router = Router()
router.message.filter(F.from_user.id == ADMIN_ID)
router.callback_query.filter(F.from_user.id == ADMIN_ID)


def admin_kb():
    today = date.today()
    kb = InlineKeyboardBuilder()
    kb.button(text="📋 Сегодня", callback_data=f"adm:day:{today.isoformat()}")
    kb.button(text="📋 Завтра", callback_data=f"adm:day:{(today + timedelta(days=1)).isoformat()}")
    kb.button(text="📊 Статистика за 30 дней", callback_data="adm:stats")
    kb.button(text="📥 Выгрузка в Excel", callback_data="adm:export")
    kb.button(text="🚫 Выходные дни", callback_data="adm:block")
    kb.adjust(2, 1, 1, 1)
    return kb.as_markup()


def back_kb():
    kb = InlineKeyboardBuilder()
    kb.button(text="← Админ-панель", callback_data="adm:home")
    return kb.as_markup()


@router.message(Command("admin"))
async def admin_home(message: Message):
    await message.answer("🛠 Админ-панель", reply_markup=admin_kb())


@router.callback_query(F.data == "adm:home")
async def admin_home_cb(cb: CallbackQuery):
    await cb.message.edit_text("🛠 Админ-панель", reply_markup=admin_kb())
    await cb.answer()


@router.callback_query(F.data.startswith("adm:day:"))
async def day_schedule(cb: CallbackQuery):
    day = cb.data.split(":", 2)[2]
    rows = db.day_bookings(day)
    if rows:
        lines = [
            f"{r['slot']} · {MASTERS[r['master']]['name']} · {SERVICES[r['service']][0]}\n"
            f"      {r['name']}, {r['phone']}"
            for r in rows
        ]
        revenue = sum(SERVICES[r["service"]][1] for r in rows)
        text = f"📋 {fmt_day(day)} — {plural(len(rows), 'запись', 'записи', 'записей')}, ≈{revenue} ₽\n\n" + "\n".join(lines)
    else:
        text = f"📋 {fmt_day(day)} — записей нет."
    await cb.message.edit_text(text, reply_markup=back_kb())
    await cb.answer()


@router.callback_query(F.data == "adm:stats")
async def stats(cb: CallbackQuery):
    today = date.today()
    rows = db.bookings_between((today - timedelta(days=30)).isoformat(), today.isoformat())
    if not rows:
        await cb.message.edit_text("За последние 30 дней записей не было.", reply_markup=back_kb())
        return await cb.answer()
    revenue = sum(SERVICES[r["service"]][1] for r in rows)
    by_master = Counter(MASTERS[r["master"]]["name"] for r in rows)
    by_service = Counter(SERVICES[r["service"]][0] for r in rows)
    ratings = [r["rating"] for r in rows if r["rating"]]
    avg = f"{sum(ratings) / len(ratings):.1f} ⭐ ({len(ratings)} оценок)" if ratings else "пока нет"
    clients = len({r["user_id"] for r in rows})
    text = (
        "📊 Статистика за 30 дней\n\n"
        f"Записей: {len(rows)}\nКлиентов: {clients}\nВыручка: ≈{revenue} ₽\n"
        f"Средний чек: ≈{revenue // len(rows)} ₽\nСредняя оценка: {avg}\n\n"
        "По мастерам:\n" + "\n".join(f"• {n}: {c}" for n, c in by_master.most_common()) +
        "\n\nПопулярные услуги:\n" + "\n".join(f"• {n}: {c}" for n, c in by_service.most_common())
    )
    await cb.message.edit_text(text, reply_markup=back_kb())
    await cb.answer()


@router.callback_query(F.data == "adm:export")
async def export(cb: CallbackQuery):
    today = date.today()
    rows = db.bookings_between((today - timedelta(days=90)).isoformat(), (today + timedelta(days=30)).isoformat())
    wb = Workbook()
    ws = wb.active
    ws.title = "Записи"
    ws.append(["Дата", "Время", "Мастер", "Услуга", "Цена, ₽", "Клиент", "Телефон", "Оценка"])
    for r in rows:
        title, price = SERVICES[r["service"]]
        ws.append([r["day"], r["slot"], MASTERS[r["master"]]["name"], title, price,
                   r["name"], r["phone"], r["rating"] or ""])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="B8741A")
    for col, width in zip("ABCDEFGH", (12, 8, 12, 20, 9, 18, 16, 8)):
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    await cb.message.answer_document(
        BufferedInputFile(buf.getvalue(), filename=f"bookings_{today.isoformat()}.xlsx"),
        caption=f"Записи за 90 дней и на месяц вперёд: {len(rows)} шт.",
    )
    await cb.answer()


def block_kb():
    blocked = db.blocked_days()
    kb = InlineKeyboardBuilder()
    for i in range(14):
        day = (date.today() + timedelta(days=i)).isoformat()
        mark = "🚫" if day in blocked else "✅"
        kb.button(text=f"{mark} {fmt_day(day)}", callback_data=f"adm:toggle:{day}")
    kb.button(text="← Админ-панель", callback_data="adm:home")
    kb.adjust(2)
    return kb.as_markup()


@router.callback_query(F.data == "adm:block")
async def block_menu(cb: CallbackQuery):
    await cb.message.edit_text("Нажмите на день, чтобы закрыть или открыть его для записи:", reply_markup=block_kb())
    await cb.answer()


@router.callback_query(F.data.startswith("adm:toggle:"))
async def toggle(cb: CallbackQuery):
    day = cb.data.split(":", 2)[2]
    closed = db.toggle_day(day)
    await cb.message.edit_reply_markup(reply_markup=block_kb())
    await cb.answer(f"{fmt_day(day)}: {'закрыт' if closed else 'открыт'} для записи")
